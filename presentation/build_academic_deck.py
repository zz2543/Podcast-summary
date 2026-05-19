from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "assets"
OUTPUT = ROOT / "GotIt-Podsum-academic-outline.pptx"
DIAGRAM = ASSET_DIR / "pipeline-architecture-ai.png"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

COLORS = {
    "ink": RGBColor(24, 31, 43),
    "muted": RGBColor(92, 104, 121),
    "line": RGBColor(214, 220, 229),
    "soft": RGBColor(247, 249, 252),
    "blue": RGBColor(28, 88, 166),
    "blue_soft": RGBColor(229, 238, 252),
    "green": RGBColor(32, 128, 96),
    "purple": RGBColor(105, 72, 172),
    "orange": RGBColor(196, 103, 31),
    "red": RGBColor(192, 51, 66),
    "gold": RGBColor(183, 136, 32),
}

TITLE_FONT = "Aptos Display"
BODY_FONT = "Aptos"
CJK_FONT = "Microsoft YaHei"


def set_font(run, size=18, bold=False, color=None):
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or COLORS["ink"]


def add_textbox(slide, x, y, w, h, text="", size=18, bold=False, color=None, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size=size, bold=bold, color=color)
    return box


def add_title(slide, title, kicker=None, slide_no=None):
    if kicker:
        add_textbox(slide, 0.65, 0.28, 10.5, 0.28, kicker, size=9, bold=True, color=COLORS["blue"])
    add_textbox(slide, 0.65, 0.55, 10.9, 0.52, title, size=23, bold=True, color=COLORS["ink"])
    if slide_no:
        add_textbox(slide, 12.18, 0.42, 0.7, 0.28, slide_no, size=8.5, color=COLORS["muted"], align=PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.16), Inches(12.0), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.fill.background()


def add_footer(slide):
    add_textbox(slide, 0.65, 7.1, 5.5, 0.2, "GotIt / Podsum 项目展示 · source: presentation/outline.md", size=7.5, color=COLORS["muted"])


def add_bullets(slide, items, x, y, w, h, size=15, color=None, gap=0.86):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = 0
        p.text = ""
        p.space_after = Pt(5 * gap)
        p.line_spacing = 1.12
        run = p.add_run()
        run.text = f"• {item}"
        set_font(run, size=size, color=color or COLORS["ink"])
    return box


def add_numbered(slide, items, x, y, w, h, size=14.5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.word_wrap = True
    for idx, item in enumerate(items, start=1):
        p = tf.paragraphs[0] if idx == 1 else tf.add_paragraph()
        p.space_after = Pt(7)
        p.line_spacing = 1.08
        run = p.add_run()
        run.text = f"{idx}. {item}"
        set_font(run, size=size, color=COLORS["ink"])
    return box


def add_card(slide, x, y, w, h, title=None, body=None, accent="blue", fill=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill or COLORS["soft"]
    shape.line.color.rgb = COLORS["line"]
    shape.line.width = Pt(1)
    if title or body:
        tf = shape.text_frame
        tf.clear()
        tf.margin_left = Inches(0.16)
        tf.margin_right = Inches(0.16)
        tf.margin_top = Inches(0.11)
        tf.margin_bottom = Inches(0.08)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.word_wrap = True
        if title:
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = title
            set_font(run, size=12.5, bold=True, color=COLORS[accent])
        if body:
            p = tf.add_paragraph() if title else tf.paragraphs[0]
            p.space_before = Pt(4)
            run = p.add_run()
            run.text = body
            set_font(run, size=11.2, color=COLORS["ink"])
    return shape


def add_placeholder(slide, x, y, w, h, label, note="后续粘贴截图 / 视频 / 图片"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(250, 251, 253)
    shape.line.color.rgb = RGBColor(184, 194, 208)
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    set_font(r, size=13.5, bold=True, color=COLORS["muted"])
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = note
    set_font(r2, size=10.5, color=COLORS["muted"])
    return shape


def add_quote(slide, x, y, w, h, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["blue_soft"]
    shape.line.color.rgb = RGBColor(181, 204, 238)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.22)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.1)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_font(r, size=16.5, bold=True, color=COLORS["blue"])
    return shape


def add_simple_table(slide, x, y, w, h, headers, rows, widths=None, font_size=9.5, highlight_rows=None):
    highlight_rows = highlight_rows or set()
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS["blue"]
        cell.text = header
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_font(r, size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, val in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 244, 244) if row_idx in highlight_rows else RGBColor(255, 255, 255)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.0
                for r in p.runs:
                    set_font(r, size=font_size, color=COLORS["red"] if row_idx in highlight_rows else COLORS["ink"])
    return table_shape


def add_diagram(slide, x, y, w, h):
    slide.shapes.add_picture(str(DIAGRAM), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_overlay_label(slide, x, y, w, h, text, size=8.5, color=None):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.fill.transparency = 15
    box.line.fill.background()
    tf = box.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_font(r, size=size, bold=True, color=color or COLORS["ink"])
    return box


def new_slide(prs, title, kicker, slide_no):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    add_title(slide, title, kicker, slide_no)
    add_footer(slide)
    return slide


def build():
    if not DIAGRAM.exists():
        raise FileNotFoundError(f"Missing diagram asset: {DIAGRAM}")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # P1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    add_textbox(slide, 0.75, 0.68, 7.4, 0.9, "GotIt", size=42, bold=True, color=COLORS["ink"])
    add_textbox(slide, 0.8, 1.62, 7.2, 0.55, "本地端到端播客摘要工具", size=23, bold=True, color=COLORS["blue"])
    add_textbox(slide, 0.82, 2.25, 7.0, 0.8, "将一小时播客压缩为 30 秒可听、可验证、可导出的结构化摘要", size=18, color=COLORS["ink"])
    add_bullets(
        slide,
        ["端到端处理：链接 / 文件输入、ASR、摘要、章节、引用、实体、TTS 与导出", "面向真实使用场景：自用工具，而非一次性 demo", "展示重点：工程完整性、技术实现与人工验证投入"],
        0.82,
        3.35,
        6.9,
        1.8,
        size=14,
    )
    add_textbox(slide, 0.82, 6.05, 4.8, 0.36, "[你的名字] · 2026.05", size=13.5, color=COLORS["muted"])
    add_placeholder(slide, 8.05, 0.82, 4.55, 5.55, "首页截图 + episode 卡片", "保留给 v2 首页或真实运行截图")
    add_footer(slide)

    # P2
    slide = new_slide(prs, "问题定义与产品定位", "P2 · Problem", "02")
    add_bullets(
        slide,
        [
            "长音频内容成本高：单集播客通常持续 40-90 分钟，筛选与回看都需要额外时间。",
            "用户目标并非完整转录，而是快速判断内容价值、结构脉络与值得回听的位置。",
            "现有工具在轻量化、章节化、摘要浓缩与原文可验证性之间仍存在明显缺口。",
            "GotIt 的目标是用单个链接触发完整流程，并输出 hook、三幕结构、章节、引用、实体与 30 秒 TTS 摘要。",
        ],
        0.8,
        1.52,
        7.2,
        2.5,
        size=14.5,
    )
    add_quote(slide, 0.8, 4.52, 7.2, 0.9, "给一个链接，10 分钟后拿到可读、可听、可验证的播客浓缩。")
    add_placeholder(slide, 8.45, 1.48, 3.9, 2.15, "1 小时音频波形")
    add_placeholder(slide, 8.45, 4.05, 3.9, 2.15, "摘要卡片 + digest 播放器")

    # P3
    slide = new_slide(prs, "Demo 导览：从输入到可听摘要", "P3 · Demo Flow", "03")
    add_numbered(
        slide,
        [
            "粘贴 YouTube / Bilibili / mp3 链接，或提交本地音频文件。",
            "通过 WebSocket 查看实时阶段进度：下载、ASR、摘要、章节、TTS。",
            "进入详情页查看 hook 引用、三幕摘要、章节时间轴与可点击原文引用。",
            "按需生成音频概览：由豆包大模型语音朗读摘要。",
            "导出 Markdown、JSON 与 MP3，便于复盘、归档或二次处理。",
        ],
        0.82,
        1.55,
        6.1,
        3.55,
        size=14.5,
    )
    add_placeholder(slide, 7.45, 1.42, 4.85, 4.3, "60 秒现场 demo / 备用视频")
    add_card(slide, 0.82, 5.55, 6.1, 0.85, "讲述策略", "本页只做口头预告，避免用截图替代现场演示；demo 翻车时切换到备用视频。", "purple")

    # P4
    slide = new_slide(prs, "端到端 Pipeline：八阶段异步处理", "P4 · Pipeline", "04")
    add_diagram(slide, 0.45, 1.28, 12.4, 5.55)
    top_labels = [
        ("输入", 0.46),
        ("fetch", 2.02),
        ("ASR", 3.62),
        ("LLM 摘要", 5.23),
        ("引用验证", 6.86),
        ("实体抽取", 8.46),
        ("导出", 10.06),
        ("TTS", 11.47),
    ]
    for text, x in top_labels:
        add_overlay_label(slide, x, 2.15, 0.88, 0.3, text, size=8.2, color=COLORS["blue"])
    add_overlay_label(slide, 2.0, 4.23, 2.2, 0.28, "React v1 / v2", size=8.2, color=COLORS["blue"])
    add_overlay_label(slide, 5.25, 4.23, 2.35, 0.28, "FastAPI · WS", size=8.2, color=COLORS["green"])
    add_overlay_label(slide, 9.05, 4.23, 2.4, 0.28, "SQLite / Event Bus", size=8.2, color=COLORS["purple"])
    add_textbox(slide, 0.72, 6.72, 11.9, 0.28, "关键机制：每个 stage 独立 retry、独立失败、独立可重做；stage_progress 写入 SQLite，前端订阅 WS 事件。", size=10.8, color=COLORS["ink"])

    # P5
    slide = new_slide(prs, "真实界面：列表页与任务状态", "P5 · UI List", "05")
    add_placeholder(slide, 0.8, 1.45, 7.25, 4.85, "列表页全屏截图")
    add_bullets(
        slide,
        [
            "状态点以低干扰方式表达运行、完成与失败状态。",
            "活动任务条显示后端实时推送的处理阶段，而不是前端轮询结果。",
            "搜索与五类状态过滤用于快速定位 episode。",
            "提交弹窗分为本地上传、Audio URL、YouTube / Bilibili 三类入口。",
            "“+ GotIt” 主按钮作为唯一高优先级操作入口。",
        ],
        8.45,
        1.48,
        3.9,
        4.25,
        size=13.2,
    )

    # P6
    slide = new_slide(prs, "真实界面：详情页与验证交互", "P6 · Detail UX", "06")
    add_placeholder(slide, 0.8, 1.42, 6.55, 4.95, "详情页滚动 GIF / 截图")
    add_numbered(
        slide,
        [
            "Hook 大字引用：用 30 字内内容定位本集主题。",
            "Sticky 原音频播放器：滚动阅读时仍可随时播放原音频。",
            "章节引用可点击：时间戳由线性插值估算，点击后 seek 到原位置并播放。",
            "AI Digest 播放器：与原音频播放器视觉区分，突出 30 秒浓缩摘要。",
        ],
        7.75,
        1.55,
        4.55,
        3.35,
        size=13.5,
    )
    add_card(slide, 7.75, 5.2, 4.55, 0.98, "设计重心", "详情页服务于“读摘要 + 回到原文验证”这一真实动作，因此交互准确性比装饰性更重要。", "green")

    # P7
    slide = new_slide(prs, "系统架构：单进程本地工具的分层实现", "P7 · Architecture", "07")
    add_diagram(slide, 0.65, 1.32, 7.2, 4.05)
    add_overlay_label(slide, 1.55, 3.0, 1.9, 0.28, "Frontend", size=8.5, color=COLORS["blue"])
    add_overlay_label(slide, 3.98, 4.2, 1.95, 0.28, "FastAPI", size=8.5, color=COLORS["green"])
    add_overlay_label(slide, 5.75, 5.47, 1.3, 0.28, "LLM / TTS", size=8.5, color=COLORS["purple"])
    add_bullets(
        slide,
        [
            "Frontend：React + CSS 的 v1 与 React + Tailwind / Framer / Aceternity 的 v2 共享后端接口。",
            "Backend：FastAPI 单进程运行，asyncio 负责 pipeline 调度，tenacity 提供重试能力。",
            "State：SQLite 存储 episode、job 与 stage_progress；文件系统保存音频、导出与中间产物。",
            "External services：Doubao ASR、DeepSeek / Qwen / Anthropic LLM、Doubao BigTTS 通过适配层接入。",
            "扩展策略：当前不上 Celery / Redis，但接口边界保留队列化与多用户化空间。",
        ],
        8.25,
        1.45,
        4.2,
        4.8,
        size=12.5,
    )

    # P8
    slide = new_slide(prs, "技术栈与选型理由", "P8 · Stack", "08")
    rows = [
        ("后端", "FastAPI / SQLAlchemy 2 / Alembic / Pydantic 2", "类型严格、async 友好、迁移可控"),
        ("任务", "asyncio / tenacity", "本地工具无需额外队列，保持实现轻量"),
        ("前端 v1", "React 19 / Vite / CSS", "功能完整，依赖最小"),
        ("前端 v2", "React 19 / Tailwind / Framer / Aceternity", "用于视觉与交互对比实验"),
        ("ASR", "豆包流式 / 录音文件大模型 2.0", "中文识别质量与成本较优"),
        ("LLM", "DeepSeek / Qwen / Anthropic", "多供应商抽象，便于成本与质量切换"),
        ("TTS", "豆包 v1 / 豆包 BigTTS 2.0", "双轨可回退，支持 30 秒音频概览"),
        ("下载", "yt-dlp / ffmpeg", "兼容主流音视频来源"),
        ("测试", "pytest / Vitest", "domain 模块覆盖率门禁"),
    ]
    add_simple_table(slide, 0.72, 1.45, 11.95, 5.25, ["层", "选型", "理由"], rows, widths=[1.25, 5.0, 5.7], font_size=9.3)

    # P9
    slide = new_slide(prs, "八阶段流水线与实时事件", "P9 · Runtime", "09")
    stages = [
        ("1", "fetch", "yt-dlp 下载与 ffmpeg 归一化"),
        ("2", "transcribe", "豆包录音文件 ASR，base64 提交与轮询查询"),
        ("3", "summarize_hook", "一句话 hook，JSON schema 强校验"),
        ("4", "summarize_three_act", "三幕结构摘要"),
        ("5", "chapter_outline", "章节与章节内 quote 候选"),
        ("6", "quote_verify", "NFKC / 空白归一 / 子串匹配 / 段内插值时间戳"),
        ("7", "entity_extract", "人名、书名、产品等命名实体"),
        ("8", "export", "Markdown 与 JSON 输出"),
        ("+", "tts", "独立触发的双向流 WebSocket 音频合成"),
    ]
    y = 1.42
    for num, name, desc in stages:
        add_card(slide, 0.78, y, 0.52, 0.42, num, None, "blue", fill=COLORS["blue_soft"])
        add_textbox(slide, 1.43, y + 0.02, 2.2, 0.35, name, size=11.5, bold=True, color=COLORS["blue"])
        add_textbox(slide, 3.45, y + 0.02, 8.8, 0.35, desc, size=11.2, color=COLORS["ink"])
        y += 0.52
    add_card(slide, 0.78, 6.32, 11.8, 0.56, "运行时事实", "每个阶段独立 retry；状态变化推送到 /api/ws/jobs；stage_progress 持久化后支持重启后继续观察。", "green")

    # P10
    slide = new_slide(prs, "多供应商抽象与 BigTTS 协议实现", "P10 · Provider Layer", "10")
    add_card(slide, 0.78, 1.45, 5.6, 4.75, "抽象层", "ASR、LLM、TTS 都通过 create_xxx_client() 工厂与 Protocol 接口接入。调用方只依赖能力协议，不感知具体供应商。", "blue")
    add_bullets(
        slide,
        [
            "ASR：doubao / openai_whisper / qwen",
            "LLM：deepseek / qwen / anthropic",
            "TTS：doubao legacy / doubao bigmodel / qwen",
            "统一切换：通过 env 配置改变默认 provider",
        ],
        1.08,
        2.55,
        4.95,
        2.45,
        size=12.5,
    )
    add_card(slide, 6.85, 1.45, 5.55, 4.75, "协议实现", "豆包 BigTTS 2.0 使用 WebSocket 二进制帧，而非一次性 HTTP 接口。项目内实现了 header、event id、payload size、session 生命周期与 AudioOnlyServer 分片接收。", "purple")
    add_placeholder(slide, 7.24, 3.78, 4.75, 1.85, "二进制帧 / factory 代码片段", "后续粘贴代码截图")

    # P11
    slide = new_slide(prs, "Hard-Won Bugs：工程投入的证据", "P11 · Debugging Evidence", "11")
    rows = [
        ("1", "TTS 401", "Authorization 需要 Bearer;<token>", "修正 header", "不能默认所有 API 遵守 OAuth 习惯"),
        ("2", "TTS 400", "v1 单次文本 ≤1024 字节", "按 UTF-8 字节切句", "协议限制常藏在脚注"),
        ("3", "BigTTS 并发超限", "resource_id 使用内部数字编号", "改为 seed-tts-2.0", "错误信息需要反证"),
        ("4", "digest.mp3 为空", "mixed 语言路由到英文音色", "mixed 归类为中文", "静默失败更危险"),
        ("5", "多个 quote 时间戳相同", "旧逻辑只返回 segment 起点", "段内字符比例插值", "要检查数据分布"),
        ("6", "Retry 后 PendingRollback", "quote_verify 重跑撞 unique 约束", "前置清理 + 测试", "幂等是基础要求"),
    ]
    add_simple_table(
        slide,
        0.55,
        1.38,
        12.25,
        4.75,
        ["#", "现象", "根因", "修法", "学到"],
        rows,
        widths=[0.45, 2.0, 3.0, 2.15, 4.65],
        font_size=8.1,
        highlight_rows={3, 5},
    )
    add_textbox(slide, 0.75, 6.82, 11.9, 0.25, "重点讲 #3 与 #5：它们最能体现协议排查、日志阅读、数据验证与测试补强。", size=10.5, color=COLORS["red"])

    # P12
    slide = new_slide(prs, "SDD：先写规范再写代码", "P12 · Spec-Driven Development", "12")
    tree = (
        "specs/001-podcast-summary/\n"
        "├── spec.md / plan.md / research.md\n"
        "├── data-model.md / tasks.md / quickstart.md\n"
        "├── ui-brief.md\n"
        "├── contracts/\n"
        "│   ├── http-api.md\n"
        "│   ├── job-events.md\n"
        "│   └── episode-output.schema.json\n"
        "└── design/...\n\n"
        ".specify/memory/constitution.md v1.0.0"
    )
    add_card(slide, 0.78, 1.45, 5.45, 4.9, "规范目录", tree, "blue")
    add_bullets(
        slide,
        [
            "先定义用户故事、验收标准、数据模型与 API 契约，再进入任务拆解与实现。",
            "Constitution 明确五条原则：英文交付、Python 3.11+、domain 覆盖率、配置外置、prompt 版本化。",
            "Domain 模块 ≥80% 覆盖率是 NON-NEGOTIABLE 质量门。",
            "所有 LLM prompt 集中维护在 prompts/，避免散落在业务代码中。",
        ],
        6.7,
        1.65,
        5.65,
        3.2,
        size=13.0,
    )
    add_placeholder(slide, 6.9, 5.18, 5.25, 1.1, "specs 目录树截图")

    # P13
    slide = new_slide(prs, "测试与质量门", "P13 · Quality Gates", "13")
    add_card(
        slide,
        0.78,
        1.45,
        5.65,
        2.2,
        "测试结构",
        "backend/tests/\n├── unit/          12 files\n└── integration/    8 files\n\nfrontend: Vitest + jsdom",
        "blue",
    )
    add_card(
        slide,
        0.78,
        3.98,
        5.65,
        1.52,
        "覆盖率门禁",
        "--cov=backend/src/podsum/domain\n--cov-fail-under=80",
        "red",
    )
    add_bullets(
        slide,
        [
            "Domain 模块尽量纯函数化，单测不依赖数据库与外部 API。",
            "quote_verifier 插值 bug 修复同步加入回归测试。",
            "流水线 stage 与 provider client 使用 mock 完成集成验证。",
            "质量门的目标不是形式覆盖，而是保护真实调试中暴露出的边界条件。",
        ],
        6.9,
        1.55,
        5.35,
        2.8,
        size=13.0,
    )
    add_placeholder(slide, 6.95, 4.75, 5.25, 1.42, "pytest --cov 终端截图")

    # P14
    slide = new_slide(prs, "两套前端：一个工程化视觉实验", "P14 · Frontend Experiment", "14")
    rows = [
        ("端口", "5173", "5174"),
        ("样式", "纯 CSS + 渐变", "Tailwind + Apple Liquid Glass"),
        ("动效", "无", "framer-motion + Aceternity"),
        ("风格", "功能优先", "内容优先 + 克制配色"),
        ("用途", "当前主力", "A/B 视觉测试"),
    ]
    add_simple_table(slide, 0.72, 1.45, 5.95, 2.25, ["维度", "v1 frontend/", "v2 frontend-v2/"], rows, widths=[1.2, 2.35, 2.4], font_size=9.0)
    add_card(slide, 0.78, 4.05, 5.8, 0.95, "共同点", "两套前端共享同一后端、API client 与 WebSocket 协议，可通过 make run-both 同时运行。", "green")
    add_placeholder(slide, 7.1, 1.45, 2.6, 4.65, "v1 截图")
    add_placeholder(slide, 9.95, 1.45, 2.6, 4.65, "v2 截图")

    # P15
    slide = new_slide(prs, "边界意识：未来路线与 Q&A", "P15 · Roadmap", "15")
    add_card(slide, 0.78, 1.45, 3.6, 3.6, "今天能做", "✓ 本地端到端跑通\n✓ 5 天 / 93 次提交\n✓ 8 stage / 20 个测试\n✓ 可导出 md / json / mp3", "green")
    add_card(slide, 4.85, 1.45, 3.6, 3.6, "下一步可做", "多用户认证\n内容审核\nPWA / 原生 App 化\n部署与成本兜底", "blue")
    add_card(slide, 8.92, 1.45, 3.6, 3.6, "有意不做", "Celery / Redis：本地工具不值得\n公开转录：版权边界不清晰\n复杂云部署：不符合当前自用定位", "orange")
    add_placeholder(slide, 0.78, 5.45, 3.6, 0.95, "Logo / GitHub 链接")
    add_textbox(slide, 5.0, 5.48, 4.0, 0.72, "Q&A", size=34, bold=True, color=COLORS["blue"], align=PP_ALIGN.CENTER)
    add_textbox(slide, 8.2, 5.55, 4.25, 0.55, "AI 是协作工具；决策、排查与验证仍由人完成。", size=13, color=COLORS["ink"])

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
